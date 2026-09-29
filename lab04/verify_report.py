# -*- coding: utf-8 -*-
import docx, zipfile, io, re, sys
f = "2025311489_赖宇轩_实验作业四.docx"
d = docx.Document(f)
z = zipfile.ZipFile(f)
fail = []
def ck(n, ok, det=""):
    print(("  PASS  " if ok else "  FAIL  ") + n + (("   " + str(det)) if det else ""))
    if not ok: fail.append(n)

txt = "\n".join(p.text for p in d.paragraphs)
print("=== 1. 模板提示语应全部删除 ===")
left = [p.text[:40] for p in d.paragraphs
        if p.text.strip().startswith("提示：")]
ck("无模板提示语残留", not left, left)

print("\n=== 2. 占位符应被填掉 ===")
ck("无 ＿＿ 下划线占位", "＿＿" not in txt)
ck("无「在此插入图片」占位",
   not any("在此插入图片" in c.text for t in d.tables for r in t.rows for c in r.cells))

print("\n=== 3. 封面 ===")
for kw in ("赖宇轩", "2025311489", "2026 年 9 月 29 日"):
    ck(f"封面含 {kw}", kw in txt)

print("\n=== 4. 11 个表逐项抽查 ===")
T = d.tables
checks = [
    (0, 1, 0, "操作系统"), (0, 2, 1, "Python 3.12.10"), (0, 4, 1, "llama.cpp"),
    (5, 1, 0, "Vector RAG"), (5, 1, 1, "1.000"), (5, 2, 5, "0.639"), (5, 3, 6, "0.639"),
    (6, 1, 2, "1.000"), (6, 2, 3, "0.690"), (6, 5, 4, "0.857"),
    (7, 1, 2, "5"), (7, 4, 2, "4"), (7, 9, 1, "Q11"),
    (8, 1, 3, "1 分钟"), (9, 1, 2, "0.500"),
    (10, 1, 1, "Q01"), (10, 4, 1, "Q17"), (10, 6, 3, "条目"),
    (11, 3, 1, "定位检索失败原因"),
]
for ti, ri, ci, want in checks:
    got = T[ti].cell(ri, ci).text
    ck(f"表[{ti}] r{ri}c{ci} 含 {want!r}", want in got, repr(got[:50]))

print("\n=== 5. 问题集 5 道自拟题 ===")
t2 = T[4]
for k in range(4, 9):
    r = t2.rows[k]
    ck(f"表2 第{k}行已填", bool(r.cells[1].text.strip() and r.cells[3].text.strip()),
       f"{r.cells[0].text} | {r.cells[3].text}")

print("\n=== 6. 忠实度表 9 行 ===")
t5 = T[7]
filled = sum(1 for k in range(1, 10) if t5.cell(k, 2).text.strip())
ck("表5 有 9 行评分", filled == 9, f"{filled} 行")

print("\n=== 7. 失败案例 6 例 ===")
t8 = T[10]
filled = sum(1 for k in range(1, 7) if t8.cell(k, 2).text.strip())
ck("表8 有 6 例", filled == 6, f"{filled} 例")

print("\n=== 8. 三张插图 ===")
media = [n for n in z.namelist() if n.startswith("word/media/")]
ck("docx 内有 3 张图片", len(media) == 3, media)

print("\n=== 9. 思考题 6 题均有作答 ===")
for i, kw in enumerate(["大模型通识课》—主讲→ 李文瀚", "1.75 / 3", "社区摘要",
                        "更新成本最低", "RAGAS", "JSON 解析"], start=1):
    ck(f"思考题({i}) 含关键内容", kw in txt, kw)

print("\n=== 10. 关键数字与实测一致 ===")
for kw in ["1.000", "0.952", "0.714", "0.478", "0.381", "0.344", "0.611", "42 tok/s"]:
    ck(f"报告含 {kw}", kw in txt or any(kw in c.text for t in T for r in t.rows for c in r.cells))

print("\n" + ("全部通过" if not fail else f"{len(fail)} 项失败: {fail}"))
sys.exit(1 if fail else 0)
